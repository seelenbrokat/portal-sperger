"""Mustererkennung / Kennzahlen aus gespeicherten Frachten."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WEEKDAYS_DE = [
    "Montag",
    "Dienstag",
    "Mittwoch",
    "Donnerstag",
    "Freitag",
    "Samstag",
    "Sonntag",
]


def _parse_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S %Z",
        "%Y-%m-%d %H:%M:%S%z",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            # %Z oft nicht parsebar – Fallback ohne Zone
            if fmt.endswith("%Z"):
                base = text.rsplit(" ", 1)[0]
                return datetime.strptime(base, "%Y-%m-%d %H:%M:%S")
            return datetime.strptime(text[: len(fmt) + 5], fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_offers_from_excel(excel_path: Path, *, route_id: str = "") -> list[dict[str, Any]]:
    if not excel_path.exists():
        return []
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    if "Frachten" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Frachten"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    offers = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        item = dict(zip(headers, row))
        item["_route_id"] = route_id
        offers.append(item)
    return offers


def load_all_offers(routes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    all_offers: list[dict[str, Any]] = []
    for route in routes:
        all_offers.extend(
            load_offers_from_excel(Path(route["excel_path"]), route_id=route["id"])
        )
    return all_offers


def compute_analytics(
    offers: list[dict[str, Any]],
    *,
    top_n: int = 8,
) -> dict[str, Any]:
    """Kennzahlen über alle übergebenen Angebote."""
    weekday_counts = Counter()
    weekday_by_field = "erstmals_gefunden"
    providers = Counter()
    destinations = Counter()
    origins = Counter()
    prices: list[float] = []
    prices_by_currency: dict[str, list[float]] = defaultdict(list)
    length_values: list[float] = []
    weight_values: list[float] = []

    for o in offers:
        dt = _parse_date(o.get(weekday_by_field)) or _parse_date(o.get("loading_earliest_date"))
        if dt:
            weekday_counts[WEEKDAYS_DE[dt.weekday()]] += 1

        anbieter = (o.get("anbieter_name") or "").strip()
        tid = o.get("anbieter_timocom_id")
        if anbieter:
            key = f"{anbieter} ({tid})" if tid not in (None, "") else anbieter
            providers[key] += 1

        dest = " ".join(
            str(x)
            for x in (
                o.get("unloading_country") or "",
                o.get("unloading_postal_code") or "",
                o.get("unloading_city") or "",
            )
            if x
        ).strip()
        if dest:
            destinations[dest] += 1

        origin = " ".join(
            str(x)
            for x in (
                o.get("loading_country") or "",
                o.get("loading_postal_code") or "",
                o.get("loading_city") or "",
            )
            if x
        ).strip()
        if origin:
            origins[origin] += 1

        amount = _num(o.get("price_amount"))
        currency = (o.get("price_currency") or "").strip() or "EUR"
        if amount is not None and amount > 0:
            prices.append(amount)
            prices_by_currency[currency].append(amount)

        length = _num(o.get("length_m"))
        if length is not None:
            length_values.append(length)
        weight = _num(o.get("weight_t"))
        if weight is not None:
            weight_values.append(weight)

    # Stärkster Wochentag in fester Reihenfolge Mo–So
    weekday_rows = [
        {"weekday": day, "count": weekday_counts.get(day, 0)} for day in WEEKDAYS_DE
    ]
    strongest = max(weekday_rows, key=lambda r: r["count"]) if offers else None
    if strongest and strongest["count"] == 0:
        strongest = None

    avg_price = round(sum(prices) / len(prices), 2) if prices else None
    price_by_currency = {
        cur: {
            "count": len(vals),
            "avg": round(sum(vals) / len(vals), 2),
            "min": round(min(vals), 2),
            "max": round(max(vals), 2),
        }
        for cur, vals in prices_by_currency.items()
        if vals
    }

    return {
        "total_offers": len(offers),
        "with_price": len(prices),
        "avg_price": avg_price,
        "price_by_currency": price_by_currency,
        "strongest_weekday": strongest,
        "weekday_counts": weekday_rows,
        "top_providers": [
            {"name": name, "count": count} for name, count in providers.most_common(top_n)
        ],
        "top_destinations": [
            {"name": name, "count": count}
            for name, count in destinations.most_common(top_n)
        ],
        "top_origins": [
            {"name": name, "count": count} for name, count in origins.most_common(top_n)
        ],
        "avg_length_m": round(sum(length_values) / len(length_values), 2)
        if length_values
        else None,
        "avg_weight_t": round(sum(weight_values) / len(weight_values), 2)
        if weight_values
        else None,
    }


def analytics_for_routes(
    routes: list[dict[str, Any]],
    *,
    filters: dict[str, Any] | None = None,
) -> dict[str, Any]:
    from filters import visible_offers

    offers = load_all_offers(routes)
    if filters is not None:
        offers = visible_offers(offers, filters)
    overall = compute_analytics(offers)
    by_route = {}
    for route in routes:
        route_offers = [o for o in offers if o.get("_route_id") == route["id"]]
        by_route[route["id"]] = {
            "title": route["title"],
            **compute_analytics(route_offers),
        }
    return {"overall": overall, "by_route": by_route, "offers": offers}
