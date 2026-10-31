"""Erzeugt HTML-Übersichtsseiten inkl. Anbieter und Routen-Navigation."""

from __future__ import annotations

import html
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from filters import (
    is_offer_hidden,
    load_filters,
    offers_for_watched,
    visible_offers,
)


DISPLAY_COLUMNS = [
    ("anbieter_timocom_id", "Timocom-ID"),
    ("anbieter_name", "Anbieter"),
    ("anbieter_ort", "Anbieter-Ort"),
    ("anbieter_telefon", "Anbieter-Telefon"),
    ("loading_country", "Von-Land"),
    ("loading_postal_code", "Von-PLZ"),
    ("loading_city", "Von"),
    ("unloading_country", "Nach-Land"),
    ("unloading_postal_code", "Nach-PLZ"),
    ("unloading_city", "Nach"),
    ("loading_earliest_date", "Laden ab"),
    ("unloading_earliest_date", "Entladen ab"),
    ("weight_t", "Gewicht (t)"),
    ("length_m", "Länge (m)"),
    ("distance_km", "km"),
    ("vehicle_body", "Aufbau"),
    ("vehicle_type", "Fahrzeugtyp"),
    ("freight_description", "Beschreibung"),
    ("price_amount", "Preis"),
    ("price_currency", "Währung"),
    ("contact_first_name", "Kontakt Vorname"),
    ("contact_last_name", "Kontakt Name"),
    ("contact_email", "Kontakt E-Mail"),
    ("contact_phone", "Kontakt Telefon"),
    ("deeplink", "Timocom-Link"),
    ("erstmals_gefunden", "Erstmals gefunden"),
    ("zuletzt_gesehen", "Zuletzt gesehen"),
]

COMMON_CSS = """
    :root {
      --bg: #f4f6f8;
      --card: #ffffff;
      --ink: #1a2332;
      --muted: #5b6775;
      --line: #d5dde5;
      --accent: #0069b4;
      --accent-soft: #e8f2fa;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: "Segoe UI", system-ui, sans-serif;
      background: var(--bg);
      color: var(--ink);
      line-height: 1.45;
    }
    header {
      background: var(--accent);
      color: #fff;
      padding: 24px 32px 0;
    }
    header h1 {
      margin: 0 0 8px;
      font-size: 1.6rem;
      font-weight: 700;
    }
    header .sub { margin: 0 0 16px; opacity: 0.92; }
    .nav {
      display: flex;
      flex-wrap: wrap;
      gap: 4px;
      margin: 0 -4px;
    }
    .nav a {
      color: #fff;
      text-decoration: none;
      font-weight: 600;
      padding: 10px 14px;
      border-radius: 8px 8px 0 0;
      opacity: 0.85;
      background: transparent;
    }
    .nav a:hover { opacity: 1; background: rgba(255,255,255,0.12); }
    .nav a.active {
      opacity: 1;
      background: var(--bg);
      color: var(--accent);
    }
    main { padding: 24px 32px 48px; }
    .stats {
      display: flex;
      gap: 16px;
      flex-wrap: wrap;
      margin-bottom: 20px;
    }
    .stat {
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 14px 18px;
      min-width: 140px;
    }
    .stat strong {
      display: block;
      font-size: 1.5rem;
      color: var(--accent);
    }
    .stat span { color: var(--muted); font-size: 0.9rem; }
    .note {
      background: var(--accent-soft);
      border-left: 4px solid var(--accent);
      padding: 12px 16px;
      margin-bottom: 20px;
      border-radius: 0 8px 8px 0;
    }
    .section { margin-bottom: 36px; }
    .section h2 {
      margin: 0 0 6px;
      font-size: 1.25rem;
      color: var(--accent);
    }
    .section .meta { margin: 0 0 12px; color: var(--muted); }
    .grid-2 {
      display: grid;
      gap: 20px;
      grid-template-columns: repeat(auto-fit, minmax(280px, 1fr));
      margin-bottom: 24px;
    }
    .panel {
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 16px 18px;
    }
    .panel h3 { margin: 0 0 12px; font-size: 1.05rem; }
    .bars { display: flex; flex-direction: column; gap: 8px; }
    .bar-row { display: grid; grid-template-columns: 100px 1fr 36px; gap: 8px; align-items: center; }
    .bar-label { color: var(--muted); font-size: 0.9rem; }
    .bar-track { background: #e8eef3; border-radius: 4px; height: 12px; overflow: hidden; }
    .bar-fill { display: block; height: 100%; background: var(--accent); }
    .bar-n { text-align: right; font-weight: 600; }
    .table-wrap {
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 8px;
      overflow: auto;
    }
    table {
      border-collapse: collapse;
      width: 100%;
      min-width: 1400px;
      font-size: 0.92rem;
    }
    th, td {
      padding: 10px 12px;
      border-bottom: 1px solid var(--line);
      text-align: left;
      vertical-align: top;
      white-space: nowrap;
    }
    th {
      position: sticky;
      top: 0;
      background: #0f3d66;
      color: #fff;
      font-weight: 600;
      z-index: 1;
    }
    tr:hover td { background: #f7fafc; }
    code {
      background: #e8eef3;
      padding: 2px 6px;
      border-radius: 4px;
      font-size: 0.9em;
    }
    a { color: var(--accent); font-weight: 600; }
    footer {
      margin-top: 16px;
      color: var(--muted);
      font-size: 0.85rem;
    }
    .downloads {
      list-style: none;
      margin: 0;
      padding: 0;
      display: grid;
      gap: 8px;
    }
    .downloads li {
      display: flex;
      flex-wrap: wrap;
      align-items: center;
      justify-content: space-between;
      gap: 10px;
      padding: 10px 12px;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: #fafcfe;
    }
    .downloads .route-name { font-weight: 600; color: var(--ink); }
    .downloads .file-name { color: var(--muted); font-size: 0.88rem; }
    .btn-dl {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 2.2rem;
      padding: 0.4rem 0.9rem;
      border-radius: 6px;
      background: var(--accent);
      color: #fff !important;
      text-decoration: none;
      font-weight: 700;
      font-size: 0.9rem;
      white-space: nowrap;
    }
    .btn-dl:hover { background: #005291; }
    .btn-dl.secondary {
      background: #fff;
      color: var(--accent) !important;
      border: 1px solid var(--accent);
    }
    .btn-dl.secondary:hover { background: var(--accent-soft); }
    .panel .btn-dl { margin-top: 10px; }
    .actions {
      display: flex;
      flex-direction: column;
      gap: 6px;
      min-width: 150px;
    }
    .actions form { margin: 0; }
    .btn-act {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      width: 100%;
      min-height: 1.9rem;
      padding: 0.25rem 0.5rem;
      border-radius: 5px;
      border: 1px solid var(--line);
      background: #fff;
      color: var(--ink);
      font-size: 0.78rem;
      font-weight: 650;
      cursor: pointer;
      white-space: nowrap;
    }
    .btn-act:hover { background: var(--accent-soft); border-color: var(--accent); }
    .btn-act.danger {
      color: #8a1f1f;
      border-color: #e2b4b4;
      background: #fff5f5;
    }
    .btn-act.danger:hover { background: #fde8e8; }
    .btn-act.primary {
      background: var(--accent);
      color: #fff;
      border-color: var(--accent);
    }
    .btn-act.primary:hover { background: #005291; }
    .flash {
      background: #e7f6ec;
      border: 1px solid #b7dfc4;
      color: #1e5c34;
      padding: 10px 14px;
      border-radius: 8px;
      margin-bottom: 16px;
    }
    .form-inline {
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
      align-items: end;
      margin: 12px 0;
    }
    .form-inline label {
      display: flex;
      flex-direction: column;
      gap: 4px;
      font-size: 0.85rem;
      color: var(--muted);
    }
    .form-inline input {
      min-width: 180px;
      padding: 8px 10px;
      border: 1px solid var(--line);
      border-radius: 6px;
      font-size: 0.95rem;
    }
    .chip-list { display: flex; flex-wrap: wrap; gap: 8px; margin: 8px 0 16px; }
    .chip {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      background: var(--accent-soft);
      border: 1px solid #c5d9ea;
      font-size: 0.9rem;
    }
    th.col-actions, td.col-actions {
      position: sticky;
      left: 0;
      z-index: 2;
      background: #eef6fc;
      min-width: 160px;
    }
    th.col-actions { background: #0a2f52 !important; color: #fff; z-index: 3; }
    th:nth-child(2), td:nth-child(2) {
      position: sticky;
      left: 160px;
      z-index: 2;
      background: #eef6fc;
    }
    th:nth-child(2) { background: #0a2f52; color: #fff; z-index: 3; }
    th:nth-child(3), td:nth-child(3) {
      position: sticky;
      left: 270px;
      z-index: 2;
      background: #eef6fc;
    }
    th:nth-child(3) { background: #0a2f52; color: #fff; z-index: 3; }
    tr:hover td.col-actions,
    tr:hover td:nth-child(2),
    tr:hover td:nth-child(3) { background: #e3f0fa; }
"""


def _load_offers_raw(excel_path: Path) -> list[dict[str, Any]]:
    if not excel_path.exists():
        return []
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb["Frachten"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    offers = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        offers.append(dict(zip(headers, row)))
    offers.sort(key=lambda o: str(o.get("zuletzt_gesehen") or ""), reverse=True)
    return offers


def _load_offers(
    excel_path: Path,
    filters: dict[str, Any] | None = None,
    *,
    include_hidden: bool = False,
) -> list[dict[str, Any]]:
    offers = _load_offers_raw(excel_path)
    if filters is None:
        return offers
    return visible_offers(offers, filters, include_hidden=include_hidden)


def _all_offers_with_route(
    routes: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for r in routes:
        for o in _load_offers_raw(Path(r["excel_path"])):
            enriched = dict(o)
            enriched["_route_id"] = r["id"]
            enriched["_route_title"] = r["title"]
            out.append(enriched)
    out.sort(key=lambda o: str(o.get("zuletzt_gesehen") or ""), reverse=True)
    return out


def _cell(o: dict[str, Any], key: str) -> str:
    val = o.get(key)
    if val is None or val == "":
        return "—"
    text = html.escape(str(val))
    if key == "deeplink" and str(val).startswith("http"):
        return f'<a href="{html.escape(str(val))}" target="_blank" rel="noopener">Öffnen</a>'
    if key == "anbieter_name":
        return f"<strong>{text}</strong>"
    if key == "anbieter_timocom_id":
        return f"<code>{text}</code>"
    return text


def _action_cell(o: dict[str, Any], *, return_to: str) -> str:
    oid = html.escape(str(o.get("offer_id") or ""))
    sid = html.escape(str(o.get("anbieter_timocom_id") or ""))
    name = html.escape(str(o.get("anbieter_name") or ""))
    ret = html.escape(return_to)
    return f"""<td class="col-actions"><div class="actions">
      <form method="post" action="/api/hide-offer">
        <input type="hidden" name="offer_id" value="{oid}" />
        <input type="hidden" name="return_to" value="{ret}" />
        <button class="btn-act" type="submit" title="Nur dieses Angebot ausblenden">Angebot ausblenden</button>
      </form>
      <form method="post" action="/api/hide-shipper"
            onsubmit="return confirm('Auftraggeber dauerhaft ausblenden?\\n{name or sid}');">
        <input type="hidden" name="timocom_id" value="{sid}" />
        <input type="hidden" name="name" value="{name}" />
        <input type="hidden" name="return_to" value="{ret}" />
        <button class="btn-act danger" type="submit" title="Alle Angebote dieses Auftraggebers dauerhaft ausblenden">Auftraggeber dauerhaft</button>
      </form>
    </div></td>"""


def _table_html(
    offers: list[dict[str, Any]],
    *,
    return_to: str = "index.html",
    show_actions: bool = True,
    extra_columns: list[tuple[str, str]] | None = None,
) -> str:
    cols = list(DISPLAY_COLUMNS)
    if extra_columns:
        cols = list(extra_columns) + cols
    header_cells = ""
    if show_actions:
        header_cells += '<th class="col-actions">Aktionen</th>'
    header_cells += "".join(f"<th>{html.escape(label)}</th>" for _, label in cols)
    col_count = len(cols) + (1 if show_actions else 0)
    if not offers:
        body = f'<tr><td colspan="{col_count}">Noch keine Treffer.</td></tr>'
    else:
        rows = []
        for o in offers:
            tds = _action_cell(o, return_to=return_to) if show_actions else ""
            tds += "".join(f"<td>{_cell(o, key)}</td>" for key, _ in cols)
            rows.append(f"<tr>{tds}</tr>")
        body = "\n".join(rows)
    return f"""<div class="table-wrap">
      <table>
        <thead><tr>{header_cells}</tr></thead>
        <tbody>
{body}
        </tbody>
      </table>
    </div>"""


def _flash_html() -> str:
    # Nach Redirect ?msg=… – rein clientseitig aus URL lesen
    return """<div id="flash" class="flash" style="display:none"></div>
<script>
(function(){
  var p = new URLSearchParams(location.search);
  var msg = p.get('msg');
  if (!msg) return;
  var el = document.getElementById('flash');
  el.textContent = msg;
  el.style.display = 'block';
  p.delete('msg');
  var q = p.toString();
  history.replaceState(null, '', location.pathname + (q ? '?' + q : '') + location.hash);
})();
</script>"""


def _nav_html(
    routes: list[dict[str, Any]],
    active_id: str | None,
    *,
    analytics_active: bool = False,
    watched_active: bool = False,
    filters_active: bool = False,
) -> str:
    items = [
        f'<a href="index.html" class="{"active" if active_id is None and not analytics_active and not watched_active and not filters_active else ""}">Alle Routen</a>',
        f'<a href="watched.html" class="{"active" if watched_active else ""}">Tisa / WOG</a>',
        f'<a href="filters.html" class="{"active" if filters_active else ""}">Ausblenden</a>',
        f'<a href="analytics.html" class="{"active" if analytics_active else ""}">Muster / Stats</a>',
    ]
    for r in routes:
        href = html.escape(Path(r["html_path"]).name)
        title = html.escape(r["title"])
        cls = "active" if r["id"] == active_id else ""
        items.append(f'<a href="{href}" class="{cls}">{title}</a>')
    return '<nav class="nav">' + "".join(items) + "</nav>"


def _ranking_table(rows: list[dict[str, Any]], name_label: str) -> str:
    if not rows:
        return "<p class='meta'>Noch keine Daten.</p>"
    body = "".join(
        f"<tr><td>{html.escape(str(r['name']))}</td><td>{r['count']}</td></tr>" for r in rows
    )
    return f"""<div class="table-wrap" style="min-width:0">
      <table style="min-width:0">
        <thead><tr><th>{html.escape(name_label)}</th><th>Anzahl</th></tr></thead>
        <tbody>{body}</tbody>
      </table>
    </div>"""


def _weekday_bars(weekday_counts: list[dict[str, Any]]) -> str:
    if not weekday_counts:
        return ""
    max_c = max((r["count"] for r in weekday_counts), default=0) or 1
    bars = []
    for r in weekday_counts:
        width = int(round(100 * r["count"] / max_c)) if max_c else 0
        bars.append(
            f"""<div class="bar-row">
          <span class="bar-label">{html.escape(r['weekday'])}</span>
          <span class="bar-track"><span class="bar-fill" style="width:{width}%"></span></span>
          <span class="bar-n">{r['count']}</span>
        </div>"""
        )
    return "<div class='bars'>" + "".join(bars) + "</div>"


def _analytics_summary_html(analytics: dict[str, Any] | None) -> str:
    if not analytics:
        return ""
    strongest = analytics.get("strongest_weekday")
    strong_txt = (
        f"{strongest['weekday']} ({strongest['count']})"
        if strongest
        else "noch unklar"
    )
    avg = analytics.get("avg_price")
    avg_txt = f"{avg} EUR" if avg is not None else "keine Preisangaben"
    with_price = analytics.get("with_price", 0)
    total = analytics.get("total_offers", 0)
    top_p = analytics.get("top_providers") or []
    top_d = analytics.get("top_destinations") or []
    top_p_txt = ", ".join(f"{p['name']} ({p['count']})" for p in top_p[:3]) or "—"
    top_d_txt = ", ".join(f"{d['name']} ({d['count']})" for d in top_d[:3]) or "—"
    return f"""
    <section class="section" id="muster">
      <h2><a href="analytics.html">Mustererkennung</a></h2>
      <p class="meta">Je mehr Suchläufe, desto stabiler die Auswertung (aktuell {total} Angebote).</p>
      <div class="stats">
        <div class="stat"><strong>{html.escape(strong_txt)}</strong><span>Stärkster Wochentag</span></div>
        <div class="stat"><strong>{html.escape(avg_txt)}</strong><span>Ø Preis ({with_price} mit Preis)</span></div>
      </div>
      <p><strong>Top-Anbieter:</strong> {html.escape(top_p_txt)}</p>
      <p><strong>Häufigste Ziele:</strong> {html.escape(top_d_txt)}</p>
      <p class="meta"><a href="analytics.html">Alle Stats öffnen →</a></p>
    </section>"""


def build_index_html(
    excel_path: Path,
    *,
    origin_label: str,
    dest_label: str,
    title: str = "Timocom Frachten",
    tz_name: str = "Europe/Berlin",
    routes: list[dict[str, Any]] | None = None,
    active_route_id: str | None = None,
    filters: dict[str, Any] | None = None,
    return_to: str | None = None,
) -> str:
    ret = return_to or (
        f"route_{active_route_id}.html" if active_route_id else "index.html"
    )
    offers = _load_offers(excel_path, filters)
    now = datetime.now(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S %Z")
    provider_count = len(
        {str(o.get("anbieter_timocom_id")) for o in offers if o.get("anbieter_timocom_id")}
    )
    hidden_n = 0
    if filters is not None:
        raw = _load_offers_raw(excel_path)
        hidden_n = sum(1 for o in raw if is_offer_hidden(o, filters))
    page_title = html.escape(title)
    nav = _nav_html(routes or [], active_route_id)

    return f"""<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <meta http-equiv="refresh" content="300" />
  <title>{page_title}</title>
  <style>{COMMON_CSS}</style>
</head>
<body>
  <header>
    <h1>{page_title}</h1>
    <p class="sub">{html.escape(origin_label)} → {html.escape(dest_label)}</p>
    {nav}
  </header>
  <main>
    {_flash_html()}
    <div class="stats">
      <div class="stat"><strong>{len(offers)}</strong><span>Treffer</span></div>
      <div class="stat"><strong>{provider_count}</strong><span>Anbieter</span></div>
      <div class="stat"><strong>{hidden_n}</strong><span>Ausgeblendet</span></div>
      <div class="stat"><strong>{html.escape(now)}</strong><span>Stand (Europe/Berlin)</span></div>
    </div>
    <div class="note">
      Angebot ausblenden = nur diese Fracht. Auftraggeber dauerhaft = alle Frachten
      dieses Anbieters. Verwaltung unter <a href="filters.html">Ausblenden</a>.
      Watchlist: <a href="watched.html">Tisa / WOG</a>.
    </div>
    {_table_html(offers, return_to=ret)}
    <footer>
      Quelle: {html.escape(str(excel_path.name))} · Seite aktualisiert sich alle 5 Minuten automatisch.
    </footer>
  </main>
</body>
</html>
"""


def write_index_html(
    excel_path: Path,
    html_path: Path,
    *,
    origin_label: str,
    dest_label: str,
    title: str = "Timocom Frachten",
    tz_name: str = "Europe/Berlin",
    routes: list[dict[str, Any]] | None = None,
    active_route_id: str | None = None,
    filters: dict[str, Any] | None = None,
) -> Path:
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(
        build_index_html(
            excel_path,
            origin_label=origin_label,
            dest_label=dest_label,
            title=title,
            tz_name=tz_name,
            routes=routes,
            active_route_id=active_route_id,
            filters=filters,
            return_to=html_path.name,
        ),
        encoding="utf-8",
    )
    return html_path


def write_hub_html(
    html_path: Path,
    routes: list[dict[str, Any]],
    *,
    tz_name: str = "Europe/Berlin",
    analytics: dict[str, Any] | None = None,
    filters: dict[str, Any] | None = None,
) -> Path:
    """Startseite: Navigation + Muster-Kurzüberblick + alle Routen-Tabellen."""
    now = datetime.now(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S %Z")
    nav = _nav_html(routes, active_id=None)
    filt = filters if filters is not None else {}

    total = 0
    hidden_total = 0
    sections = []
    for r in routes:
        raw = _load_offers_raw(Path(r["excel_path"]))
        offers = visible_offers(raw, filt)
        hidden_total += len(raw) - len(offers)
        total += len(offers)
        href = html.escape(Path(r["html_path"]).name)
        sections.append(
            f"""
    <section class="section" id="{html.escape(r['id'])}">
      <h2><a href="{href}">{html.escape(r['title'])}</a> · {len(offers)} Treffer</h2>
      <p class="meta">{html.escape(r['origin_label'])} → {html.escape(r['dest_label'])}</p>
      {_table_html(offers, return_to='index.html')}
    </section>"""
        )

    route_titles = " · ".join(html.escape(r["title"]) for r in routes)
    content = f"""<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <meta http-equiv="refresh" content="300" />
  <title>Timocom Frachten – Alle Routen</title>
  <style>{COMMON_CSS}</style>
</head>
<body>
  <header>
    <h1>Timocom Frachten – Alle Routen</h1>
    <p class="sub">{route_titles}</p>
    {nav}
  </header>
  <main>
    {_flash_html()}
    <div class="stats">
      <div class="stat"><strong>{total}</strong><span>Treffer gesamt</span></div>
      <div class="stat"><strong>{len(routes)}</strong><span>Routen</span></div>
      <div class="stat"><strong>{hidden_total}</strong><span>Ausgeblendet</span></div>
      <div class="stat"><strong>{html.escape(now)}</strong><span>Stand (Europe/Berlin)</span></div>
    </div>
    {_analytics_summary_html(analytics)}
{"".join(sections)}
    <footer>Seite aktualisiert sich alle 5 Minuten automatisch.</footer>
  </main>
</body>
</html>
"""
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(content, encoding="utf-8")
    return html_path


def write_analytics_html(
    html_path: Path,
    routes: list[dict[str, Any]],
    stats: dict[str, Any],
    *,
    tz_name: str = "Europe/Berlin",
) -> Path:
    """Eigene Seite für Mustererkennung / Kennzahlen."""
    now = datetime.now(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S %Z")
    nav = _nav_html(routes, active_id=None, analytics_active=True)
    overall = stats.get("overall") or {}
    strongest = overall.get("strongest_weekday")
    strong_txt = (
        f"{strongest['weekday']} ({strongest['count']} Treffer)"
        if strongest
        else "noch unklar – mehr Historie nötig"
    )
    avg = overall.get("avg_price")
    avg_txt = f"{avg}" if avg is not None else "—"
    price_note = ""
    for cur, info in (overall.get("price_by_currency") or {}).items():
        price_note += (
            f"<li>{html.escape(cur)}: Ø {info['avg']} "
            f"(n={info['count']}, min {info['min']}, max {info['max']})</li>"
        )
    if not price_note:
        price_note = "<li>Keine Preise in den aktuellen Treffern.</li>"

    by_route_blocks = []
    download_items = []
    for route in routes:
        rstats = (stats.get("by_route") or {}).get(route["id"]) or {}
        r_strong = rstats.get("strongest_weekday")
        r_strong_txt = (
            f"{r_strong['weekday']} ({r_strong['count']})" if r_strong else "—"
        )
        excel_name = Path(route["excel_path"]).name
        excel_href = html.escape(excel_name)
        excel_exists = Path(route["excel_path"]).exists()
        dl_btn = (
            f'<a class="btn-dl" href="{excel_href}" download="{excel_href}">Excel laden</a>'
            if excel_exists
            else '<span class="meta">Excel noch nicht vorhanden</span>'
        )
        by_route_blocks.append(
            f"""
      <div class="panel">
        <h3>{html.escape(route['title'])}</h3>
        <p class="meta">{rstats.get('total_offers', 0)} Angebote ·
           stärkster Tag: {html.escape(r_strong_txt)} ·
           Ø Preis: {rstats.get('avg_price') if rstats.get('avg_price') is not None else '—'}</p>
        {_weekday_bars(rstats.get('weekday_counts') or [])}
        {dl_btn}
      </div>"""
        )
        download_items.append(
            f"""<li>
          <div>
            <div class="route-name">{html.escape(route['title'])}</div>
            <div class="file-name">{excel_href}</div>
          </div>
          {dl_btn if excel_exists else '<span class="meta">fehlt</span>'}
        </li>"""
        )

    downloads_html = (
        "<ul class='downloads'>" + "".join(download_items) + "</ul>"
        if download_items
        else "<p class='meta'>Noch keine Excel-Dateien.</p>"
    )

    content = f"""<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <meta http-equiv="refresh" content="300" />
  <title>Timocom Frachten – Muster / Stats</title>
  <style>{COMMON_CSS}</style>
</head>
<body>
  <header>
    <h1>Mustererkennung / Stats</h1>
    <p class="sub">Wochentage, Preise, Top-Anbieter, häufigste Ziele – Stand {html.escape(now)}</p>
    {nav}
  </header>
  <main>
    <div class="note">
      Auswertung über alle gespeicherten Angebote. Mit der Zeit (viele Suchläufe)
      werden Muster wie der stärkste Wochentag aussagekräftiger.
    </div>
    <div class="stats">
      <div class="stat"><strong>{overall.get('total_offers', 0)}</strong><span>Angebote gesamt</span></div>
      <div class="stat"><strong>{html.escape(strong_txt)}</strong><span>Stärkster Wochentag</span></div>
      <div class="stat"><strong>{html.escape(str(avg_txt))}</strong><span>Ø Preis</span></div>
      <div class="stat"><strong>{overall.get('avg_length_m') if overall.get('avg_length_m') is not None else '—'}</strong><span>Ø Lademeter</span></div>
      <div class="stat"><strong>{overall.get('avg_weight_t') if overall.get('avg_weight_t') is not None else '—'}</strong><span>Ø Gewicht (t)</span></div>
    </div>

    <section class="section" id="excel-download">
      <h2>Excel herunterladen</h2>
      <p class="meta">Vollständige Listen inkl. Anbieter und Historie – eine Datei pro Route.</p>
      {downloads_html}
    </section>

    <div class="grid-2">
      <div class="panel">
        <h3>Treffer nach Wochentag (erstmals gefunden)</h3>
        {_weekday_bars(overall.get('weekday_counts') or [])}
      </div>
      <div class="panel">
        <h3>Preise</h3>
        <ul>{price_note}</ul>
        <p class="meta">Nur Angebote mit Preisangabe ({overall.get('with_price', 0)}).</p>
      </div>
      <div class="panel">
        <h3>Top-Anbieter</h3>
        {_ranking_table(overall.get('top_providers') or [], "Anbieter")}
      </div>
      <div class="panel">
        <h3>Häufigste Destinationen</h3>
        {_ranking_table(overall.get('top_destinations') or [], "Ziel")}
      </div>
      <div class="panel">
        <h3>Häufigste Abholorte</h3>
        {_ranking_table(overall.get('top_origins') or [], "Start")}
      </div>
    </div>

    <section class="section">
      <h2>Pro Route</h2>
      <div class="grid-2">
{"".join(by_route_blocks)}
      </div>
    </section>
    <footer>Quelle: Excel-Dateien der Routen · aktualisiert nach jedem Suchlauf</footer>
  </main>
</body>
</html>
"""
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(content, encoding="utf-8")
    return html_path


def write_watched_html(
    html_path: Path,
    routes: list[dict[str, Any]],
    *,
    filters: dict[str, Any],
    tz_name: str = "Europe/Berlin",
) -> Path:
    """Alle Frachten von Watchlist-Firmen (Standard: Tisa, WOG Logistics)."""
    now = datetime.now(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S %Z")
    nav = _nav_html(routes, active_id=None, watched_active=True)
    all_offers = _all_offers_with_route(routes)
    watched = filters.get("watched_shippers") or []
    matched = offers_for_watched(all_offers, filters)

    chips = "".join(
        f'<span class="chip"><strong>{html.escape(w["label"])}</strong>'
        f'<span class="meta">Match: „{html.escape(w["match"])}“</span></span>'
        for w in watched
    ) or '<span class="meta">Keine Watchlist-Firmen.</span>'

    sections = []
    for w in watched:
        subset = offers_for_watched(all_offers, filters, watch_label=w["label"])
        for o in subset:
            o["route_title"] = o.get("_route_title") or ""
        sections.append(
            f"""
    <section class="section" id="watch-{html.escape(w['label'])}">
      <h2>{html.escape(w['label'])} · {len(subset)} Frachten</h2>
      <p class="meta">Namensfilter (enthält): <code>{html.escape(w['match'])}</code>
         · über alle Routen · ausgeblendete Auftraggeber/Angebote fehlen hier</p>
      {_table_html(
          subset,
          return_to='watched.html',
          extra_columns=[('route_title', 'Route'), ('_watch_label', 'Watch')],
      )}
    </section>"""
        )

    for o in matched:
        o["route_title"] = o.get("_route_title") or ""

    content = f"""<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <meta http-equiv="refresh" content="300" />
  <title>Timocom – Tisa / WOG Logistics</title>
  <style>{COMMON_CSS}</style>
</head>
<body>
  <header>
    <h1>Watchlist: Tisa / WOG Logistics</h1>
    <p class="sub">Frachten dieser Firmen über alle Routen – Stand {html.escape(now)}</p>
    {nav}
  </header>
  <main>
    {_flash_html()}
    <div class="note">
      Zeigt alle gespeicherten Angebote, deren Anbietername den Match-Text enthält
      (Groß/Klein egal). Firmen unter <a href="filters.html">Ausblenden</a> erweiterbar.
    </div>
    <div class="stats">
      <div class="stat"><strong>{len(matched)}</strong><span>Treffer Watchlist</span></div>
      <div class="stat"><strong>{len(watched)}</strong><span>Firmen</span></div>
    </div>
    <div class="chip-list">{chips}</div>
{"".join(sections) if sections else "<p class='meta'>Keine Treffer.</p>"}
    <footer>Daten aus Excel-Historie · Filter dauerhaft in <code>filters.json</code></footer>
  </main>
</body>
</html>
"""
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(content, encoding="utf-8")
    return html_path


def write_filters_html(
    html_path: Path,
    routes: list[dict[str, Any]],
    *,
    filters: dict[str, Any],
    tz_name: str = "Europe/Berlin",
) -> Path:
    """Verwaltung: ausgeblendete Angebote/Auftraggeber + Watchlist."""
    now = datetime.now(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S %Z")
    nav = _nav_html(routes, active_id=None, filters_active=True)
    all_offers = _all_offers_with_route(routes)
    by_id = {str(o.get("offer_id")): o for o in all_offers if o.get("offer_id")}

    hidden_offer_rows = []
    for oid in filters.get("hidden_offers") or []:
        o = by_id.get(str(oid), {})
        name = html.escape(str(o.get("anbieter_name") or "—"))
        route = html.escape(str(o.get("_route_title") or "—"))
        hidden_offer_rows.append(
            f"""<tr>
              <td><code>{html.escape(str(oid))}</code></td>
              <td>{name}</td>
              <td>{route}</td>
              <td>
                <form method="post" action="/api/unhide-offer">
                  <input type="hidden" name="offer_id" value="{html.escape(str(oid))}" />
                  <input type="hidden" name="return_to" value="filters.html" />
                  <button class="btn-act primary" type="submit">Wieder einblenden</button>
                </form>
              </td>
            </tr>"""
        )
    if not hidden_offer_rows:
        hidden_offer_rows.append(
            '<tr><td colspan="4">Keine einzelnen Angebote ausgeblendet.</td></tr>'
        )

    shipper_rows = []
    for s in filters.get("hidden_shippers") or []:
        sid = html.escape(str(s.get("timocom_id") or ""))
        sname = html.escape(str(s.get("name") or ""))
        added = html.escape(str(s.get("added_at") or "—"))
        shipper_rows.append(
            f"""<tr>
              <td><code>{sid or "—"}</code></td>
              <td><strong>{sname or "—"}</strong></td>
              <td>{added}</td>
              <td>
                <form method="post" action="/api/unhide-shipper">
                  <input type="hidden" name="timocom_id" value="{sid}" />
                  <input type="hidden" name="name" value="{sname}" />
                  <input type="hidden" name="return_to" value="filters.html" />
                  <button class="btn-act primary" type="submit">Freigeben</button>
                </form>
              </td>
            </tr>"""
        )
    if not shipper_rows:
        shipper_rows.append(
            '<tr><td colspan="4">Keine Auftraggeber dauerhaft ausgeblendet.</td></tr>'
        )

    watch_rows = []
    for w in filters.get("watched_shippers") or []:
        lab = html.escape(str(w.get("label") or ""))
        match = html.escape(str(w.get("match") or ""))
        watch_rows.append(
            f"""<tr>
              <td><strong>{lab}</strong></td>
              <td><code>{match}</code></td>
              <td>
                <form method="post" action="/api/watch-remove"
                      onsubmit="return confirm('„{lab}“ von Watchlist entfernen?');">
                  <input type="hidden" name="label" value="{lab}" />
                  <input type="hidden" name="return_to" value="filters.html" />
                  <button class="btn-act danger" type="submit">Entfernen</button>
                </form>
              </td>
            </tr>"""
        )

    content = f"""<!DOCTYPE html>
<html lang="de">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Timocom – Ausblenden / Watchlist</title>
  <style>{COMMON_CSS}</style>
</head>
<body>
  <header>
    <h1>Ausblenden &amp; Watchlist</h1>
    <p class="sub">Dauerhafte Filter · Stand {html.escape(now)}</p>
    {nav}
  </header>
  <main>
    {_flash_html()}
    <div class="note">
      Ausgeblendete Angebote und Auftraggeber erscheinen nicht mehr in den Tabellen
      und der Watchlist. Excel behält die Rohdaten. Filter liegen in
      <code>data/filters.json</code> auf dem Server.
    </div>

    <section class="section">
      <h2>Auftraggeber dauerhaft ausblenden</h2>
      <p class="meta">Per Name und/oder Timocom-ID. Wirkt auf alle Routen.</p>
      <form class="form-inline" method="post" action="/api/hide-shipper">
        <label>Timocom-ID
          <input name="timocom_id" placeholder="z. B. 123456" />
        </label>
        <label>Name
          <input name="name" placeholder="Firmenname (oder Timocom-ID)" />
        </label>
        <label>Notiz
          <input name="note" placeholder="optional" />
        </label>
        <input type="hidden" name="return_to" value="filters.html" />
        <button class="btn-act danger" type="submit">Dauerhaft ausblenden</button>
      </form>
      <div class="table-wrap" style="min-width:0">
        <table style="min-width:0">
          <thead><tr><th>Timocom-ID</th><th>Name</th><th>Seit</th><th></th></tr></thead>
          <tbody>
{"".join(shipper_rows)}
          </tbody>
        </table>
      </div>
    </section>

    <section class="section">
      <h2>Einzelne Angebote ausgeblendet</h2>
      <div class="table-wrap" style="min-width:0">
        <table style="min-width:0">
          <thead><tr><th>Angebot-ID</th><th>Anbieter</th><th>Route</th><th></th></tr></thead>
          <tbody>
{"".join(hidden_offer_rows)}
          </tbody>
        </table>
      </div>
    </section>

    <section class="section">
      <h2>Watchlist-Firmen (Tisa / WOG …)</h2>
      <p class="meta">Match = Teilstring im Anbieternamen (ohne Groß/Klein).</p>
      <form class="form-inline" method="post" action="/api/watch-add">
        <label>Anzeigename
          <input name="label" placeholder="z. B. Tisa" required />
        </label>
        <label>Match-Text
          <input name="match" placeholder="z. B. tisa" />
        </label>
        <input type="hidden" name="return_to" value="filters.html" />
        <button class="btn-act primary" type="submit">Zur Watchlist</button>
      </form>
      <div class="table-wrap" style="min-width:0">
        <table style="min-width:0">
          <thead><tr><th>Label</th><th>Match</th><th></th></tr></thead>
          <tbody>
{"".join(watch_rows) if watch_rows else '<tr><td colspan="3">Leer</td></tr>'}
          </tbody>
        </table>
      </div>
      <p class="meta"><a href="watched.html">Watchlist-Frachten öffnen →</a></p>
    </section>
  </main>
</body>
</html>
"""
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(content, encoding="utf-8")
    return html_path
