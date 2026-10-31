"""Persistiert Frachtenangebote in einer Excel-Tabelle (alle Details)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "offer_id",
    "erstmals_gefunden",
    "zuletzt_gesehen",
    "suchlauf",
    "freight_description",
    "public_remark",
    "internal_remark",
    "weight_t",
    "length_m",
    "distance_km",
    "price_amount",
    "price_currency",
    "payment_due_within_days",
    "accept_quotes",
    "trackable",
    "use_messenger",
    "creation_datetime",
    "vehicle_body",
    "vehicle_type",
    "vehicle_body_property",
    "vehicle_equipment",
    "vehicle_load_securing",
    "vehicle_swap_body",
    "additional_information",
    "logistics_document_types",
    "loading_city",
    "loading_postal_code",
    "loading_country",
    "loading_lat",
    "loading_lon",
    "loading_earliest_date",
    "loading_latest_date",
    "loading_start_time",
    "loading_end_time",
    "unloading_city",
    "unloading_postal_code",
    "unloading_country",
    "unloading_lat",
    "unloading_lon",
    "unloading_earliest_date",
    "unloading_latest_date",
    "unloading_start_time",
    "unloading_end_time",
    "loading_places_json",
    "contact_first_name",
    "contact_last_name",
    "contact_email",
    "contact_phone",
    "contact_mobile",
    "anbieter_timocom_id",
    "anbieter_name",
    "anbieter_telefon",
    "anbieter_land",
    "anbieter_plz",
    "anbieter_ort",
    "anbieter_strasse",
    "anbieter_ust_id",
    "deeplink",
    "raw_json",
]

HEADER_FILL = PatternFill("solid", fgColor="0069B4")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _join(values: Any) -> str:
    if not values:
        return ""
    if isinstance(values, list):
        return ", ".join(str(v) for v in values if v is not None)
    return str(values)


def _place(places: list[dict], loading_type: str) -> dict:
    for p in places or []:
        if p.get("loadingType") == loading_type:
            return p
    return {}


def flatten_offer(offer: dict[str, Any], *, now_local: str, run_id: str) -> dict[str, Any]:
    places = offer.get("loadingPlaces") or []
    loading = _place(places, "LOADING")
    unloading = _place(places, "UNLOADING")
    la = loading.get("address") or {}
    ua = unloading.get("address") or {}
    lg = la.get("geoCoordinate") or {}
    ug = ua.get("geoCoordinate") or {}
    vp = offer.get("vehicleProperties") or {}
    price = offer.get("price") or {}
    contact = offer.get("contactPerson") or {}
    customer = offer.get("customer") or {}

    import json

    return {
        "offer_id": offer.get("id") or "",
        "erstmals_gefunden": now_local,
        "zuletzt_gesehen": now_local,
        "suchlauf": run_id,
        "freight_description": offer.get("freightDescription") or "",
        "public_remark": offer.get("publicRemark") or "",
        "internal_remark": offer.get("internalRemark") or "",
        "weight_t": offer.get("weight_t"),
        "length_m": offer.get("length_m"),
        "distance_km": offer.get("distance_km"),
        "price_amount": price.get("amount") if isinstance(price, dict) else "",
        "price_currency": price.get("currency") if isinstance(price, dict) else "",
        "payment_due_within_days": offer.get("paymentDueWithinDays"),
        "accept_quotes": offer.get("acceptQuotes"),
        "trackable": offer.get("trackable"),
        "use_messenger": offer.get("useMessenger"),
        "creation_datetime": offer.get("creationDateTime") or "",
        "vehicle_body": _join(vp.get("body")),
        "vehicle_type": _join(vp.get("type")),
        "vehicle_body_property": _join(vp.get("bodyProperty")),
        "vehicle_equipment": _join(vp.get("equipment")),
        "vehicle_load_securing": _join(vp.get("loadSecuring")),
        "vehicle_swap_body": _join(vp.get("swapBody")),
        "additional_information": _join(offer.get("additionalInformation")),
        "logistics_document_types": _join(offer.get("logisticsDocumentTypes")),
        "loading_city": la.get("city") or "",
        "loading_postal_code": la.get("postalCode") or "",
        "loading_country": la.get("country") or "",
        "loading_lat": lg.get("latitude"),
        "loading_lon": lg.get("longitude"),
        "loading_earliest_date": loading.get("earliestLoadingDate") or "",
        "loading_latest_date": loading.get("latestLoadingDate") or "",
        "loading_start_time": loading.get("startTime") or "",
        "loading_end_time": loading.get("endTime") or "",
        "unloading_city": ua.get("city") or "",
        "unloading_postal_code": ua.get("postalCode") or "",
        "unloading_country": ua.get("country") or "",
        "unloading_lat": ug.get("latitude"),
        "unloading_lon": ug.get("longitude"),
        "unloading_earliest_date": unloading.get("earliestLoadingDate") or "",
        "unloading_latest_date": unloading.get("latestLoadingDate") or "",
        "unloading_start_time": unloading.get("startTime") or "",
        "unloading_end_time": unloading.get("endTime") or "",
        "loading_places_json": json.dumps(places, ensure_ascii=False),
        "contact_first_name": contact.get("firstName") or "",
        "contact_last_name": contact.get("lastName") or "",
        "contact_email": contact.get("email") or "",
        "contact_phone": contact.get("businessPhone") or "",
        "contact_mobile": contact.get("mobilePhone") or "",
        "anbieter_timocom_id": customer.get("id") if isinstance(customer, dict) else "",
        "anbieter_name": customer.get("name") if isinstance(customer, dict) else "",
        "anbieter_telefon": customer.get("phone") if isinstance(customer, dict) else "",
        "anbieter_land": (customer.get("companyAddress") or {}).get("country")
        if isinstance(customer, dict)
        else "",
        "anbieter_plz": (customer.get("companyAddress") or {}).get("postalCode")
        if isinstance(customer, dict)
        else "",
        "anbieter_ort": (customer.get("companyAddress") or {}).get("city")
        if isinstance(customer, dict)
        else "",
        "anbieter_strasse": (customer.get("companyAddress") or {}).get("streetOrPostbox")
        if isinstance(customer, dict)
        else "",
        "anbieter_ust_id": customer.get("taxId") if isinstance(customer, dict) else "",
        "deeplink": offer.get("deeplink") or "",
        "raw_json": json.dumps(offer, ensure_ascii=False),
    }


class ExcelStore:
    def __init__(self, path: Path, tz: ZoneInfo):
        self.path = path
        self.tz = tz
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            self._create()

    def _create(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Frachten"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for i, _ in enumerate(HEADERS, 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
        ws.column_dimensions["E"].width = 40
        ws.column_dimensions["AW"].width = 50
        ws.freeze_panes = "A2"

        log = wb.create_sheet("Suchlaeufe")
        log.append(["lauf_id", "zeitpunkt", "treffer", "neu", "aktualisiert", "fehler"])
        for cell in log[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        wb.save(self.path)

    def upsert_offers(self, offers: list[dict[str, Any]], run_id: str) -> tuple[int, int]:
        """Schreibt/aktualisiert Angebote. Returns (neu, aktualisiert)."""
        now_local = datetime.now(self.tz).strftime("%Y-%m-%d %H:%M:%S %Z")
        wb = load_workbook(self.path)
        ws = wb["Frachten"]
        # Index: offer_id -> row
        id_col = 1
        existing: dict[str, int] = {}
        for row in range(2, ws.max_row + 1):
            oid = ws.cell(row, id_col).value
            if oid:
                existing[str(oid)] = row

        new_count = 0
        updated_count = 0
        for offer in offers:
            flat = flatten_offer(offer, now_local=now_local, run_id=run_id)
            oid = str(flat["offer_id"])
            if not oid:
                continue
            if oid in existing:
                row = existing[oid]
                # erstmals_gefunden behalten
                first_seen = ws.cell(row, HEADERS.index("erstmals_gefunden") + 1).value
                values = [flat[h] for h in HEADERS]
                values[HEADERS.index("erstmals_gefunden")] = first_seen or flat["erstmals_gefunden"]
                for col, val in enumerate(values, 1):
                    ws.cell(row, col, val)
                updated_count += 1
            else:
                ws.append([flat[h] for h in HEADERS])
                existing[oid] = ws.max_row
                new_count += 1

        wb.save(self.path)
        return new_count, updated_count

    def log_run(
        self,
        run_id: str,
        treffer: int,
        neu: int,
        aktualisiert: int,
        fehler: str = "",
    ) -> None:
        wb = load_workbook(self.path)
        log = wb["Suchlaeufe"]
        now_local = datetime.now(self.tz).strftime("%Y-%m-%d %H:%M:%S %Z")
        log.append([run_id, now_local, treffer, neu, aktualisiert, fehler])
        wb.save(self.path)

    def offers_found_on_date(self, day: str) -> list[dict[str, Any]]:
        """Alle Zeilen, deren erstmals_gefunden mit day (YYYY-MM-DD) beginnt."""
        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb["Frachten"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h) for h in rows[0]]
        result = []
        idx_first = headers.index("erstmals_gefunden")
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            first = str(row[idx_first] or "")
            if first.startswith(day):
                result.append(dict(zip(headers, row)))
        return result

    def all_offers_seen_on_date(self, day: str) -> list[dict[str, Any]]:
        """Zeilen, die an day erstmals gefunden oder zuletzt gesehen wurden."""
        wb = load_workbook(self.path, read_only=True, data_only=True)
        ws = wb["Frachten"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h) for h in rows[0]]
        i_first = headers.index("erstmals_gefunden")
        i_last = headers.index("zuletzt_gesehen")
        result = []
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            first = str(row[i_first] or "")
            last = str(row[i_last] or "")
            if first.startswith(day) or last.startswith(day):
                result.append(dict(zip(headers, row)))
        return result
